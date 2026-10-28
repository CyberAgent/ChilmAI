"""リファレンスアプリ（Web UI・サーバ API）の差し替え口。

自治体ごとに異なる前処理・入力ファイルの読み込み・結果 Excel・画面の入力項目を、
``apps/`` を改変せずに差し替えるための仕組み。``WebCustomization`` を継承した
クラスを ``apps/`` の外に実装し、環境変数 ``CHILMAI_WEB_CUSTOMIZATION`` に
``module:ClassName`` 形式で指定する::

    CHILMAI_WEB_CUSTOMIZATION=mytown.web:MyTownCustomization \\
        uv run uvicorn apps.api.main:app --host 127.0.0.1 --port 8501

実装例は ``examples/sample_web_customization.py`` を参照。
"""

from __future__ import annotations

import importlib
import os
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any, ClassVar, Literal

import openpyxl

from chilmai.generic.config import SCHEMA_VERSION
from chilmai.generic.error_codes import ChilmError, ErrorCode
from chilmai.generic.parser import InputParser
from chilmai.generic.preprocessor import BasePreprocessor

ENV_VAR = "CHILMAI_WEB_CUSTOMIZATION"

_PARAM_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


@dataclass(frozen=True)
class ParamField:
    """Web UI のアップロード画面に追加する入力項目。

    入力値は文字列として ``WebCustomization`` の各メソッドに ``params`` で渡される。
    ``checkbox`` はチェックありで ``"true"``、なしで ``"false"`` になる。

    Attributes:
        name: 項目名。英数字と ``_`` のみ。``params`` のキーになる。
        label: 画面に表示する項目名。
        type: 入力の種類。``"text"``・``"number"``・``"select"``・``"checkbox"``。
        choices: ``select`` の選択肢。
        default: 初期値。未入力の場合もこの値を使う。
        required: 入力必須かどうか。
        help_text: 入力欄の下に表示する説明。
    """

    name: str
    label: str
    type: Literal["text", "number", "select", "checkbox"] = "text"
    choices: tuple[str, ...] = ()
    default: str = ""
    required: bool = False
    help_text: str = ""


class WebCustomization:
    """リファレンスアプリのカスタマイズの基底クラス。デフォルト実装は標準の挙動。

    必要なメソッド・クラス属性だけをオーバーライドする。``params`` には
    ``fields`` で定義した項目の入力値が ``{name: 値}`` で入る。
    """

    #: アップロード画面に追加する入力項目。
    fields: ClassVar[Sequence[ParamField]] = ()

    def create_preprocessor(self, params: Mapping[str, str]) -> BasePreprocessor | None:
        """データ確認・マッチングで使う前処理フックを返す。

        リクエストごとに呼ばれる。``None`` を返すと前処理なしで実行する。
        """
        return None

    def create_parser(self) -> InputParser:
        """入力ファイルの読み込みに使う ``InputParser`` を返す。

        対象のシートや列名の行を変える場合は、``SHEET_NAME``・``HEADER_ROW`` を
        上書きした ``InputParser`` のサブクラスを返す。読み込み処理そのものを
        変える場合は ``InputParser.read_table`` をオーバーライドする。
        項目名の設定画面の列名の自動検出にも使われる。
        """
        return InputParser()

    def build_result_workbook(
        self,
        result: dict[str, Any],
        meta: Mapping[str, str],
        params: Mapping[str, str],
    ) -> openpyxl.Workbook:
        """Web UI からダウンロードする結果 Excel を作成する。

        ``=`` で始まる文字列が数式として実行されないよう、戻り値のすべてのセルは
        文字列として保存される。そのため、数式は結果 Excel に出力できない。

        Args:
            result: ``MatchingService.match()`` の戻り値。``output_columns``・
                ``output_rows`` に結果ファイルの列・行が入る。
            meta: 実行時の情報。``active_profile``（プロファイル名）と
                ``executed_at``（実行日時）。
            params: 追加の入力項目の値。
        """
        return build_default_workbook(result, meta, params)

    def result_filename(self, params: Mapping[str, str]) -> str:
        """結果 Excel のダウンロード時のファイル名を返す。"""
        return "matching_result.xlsx"


def build_default_workbook(
    result: dict[str, Any],
    meta: Mapping[str, str] | None = None,
    params: Mapping[str, str] | None = None,
) -> openpyxl.Workbook:
    """標準の結果 Excel（``result``・``meta`` の 2 シート）を作成する。

    ``WebCustomization.build_result_workbook`` で、この結果にシートや書式を
    追加する場合にも使える。
    """
    output_columns: list[str] = result.get("output_columns", [])
    output_rows: list[dict[str, Any]] = result.get("output_rows", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "result"
    ws.append(output_columns)
    for row in output_rows:
        ws.append([row.get(col, "") for col in output_columns])

    meta_ws = wb.create_sheet(title="meta")
    meta_ws.append(["key", "value"])
    meta_items: list[tuple[str, str]] = [
        ("schema_version", str(SCHEMA_VERSION)),
        ("active_profile", (meta or {}).get("active_profile", "")),
        ("executed_at", (meta or {}).get("executed_at", "")),
    ]
    meta_items.extend((f"param.{name}", value) for name, value in (params or {}).items())
    for key, value in meta_items:
        meta_ws.append([key, value])
    return wb


def neutralize_formulas(wb: openpyxl.Workbook) -> None:
    """すべてのシートの数式セルを文字列セルに変える。

    openpyxl は ``=`` で始まる文字列を数式セルとして扱う。入力値に含まれる文字列が
    Excel で数式として実行されないよう、文字列型に戻す。
    """
    for ws in wb.worksheets:
        for ws_row in ws.iter_rows():
            for cell in ws_row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def _invalid_param(message: str) -> ChilmError:
    return ChilmError(message, code=ErrorCode.INVALID_INPUT_PARAMETER)


def parse_params(fields: Sequence[ParamField], raw: Mapping[str, Any]) -> dict[str, str]:
    """入力値を検証し、``fields`` の定義どおりの ``{name: 値}`` に整える。

    Raises:
        ChilmError: 定義にない項目がある場合、必須項目が未入力の場合、
            選択肢にない値や数値でない値が入力された場合。
    """
    known = {f.name for f in fields}
    unknown = sorted(str(key) for key in raw if key not in known)
    if unknown:
        raise _invalid_param(f"定義されていない入力項目が指定されました: {unknown}")

    params: dict[str, str] = {}
    for field in fields:
        value = raw.get(field.name)
        if field.type == "checkbox":
            if value is None:
                text = field.default or "false"
            else:
                text = "true" if str(value).strip().lower() in ("true", "on", "1") else "false"
            params[field.name] = text
            continue
        text = str(value).strip() if value is not None else ""
        if not text:
            text = field.default
        if not text:
            if field.required:
                raise _invalid_param(f"「{field.label}」を入力してください。")
            params[field.name] = ""
            continue
        if field.type == "select" and text not in field.choices:
            raise _invalid_param(f"「{field.label}」には選択肢の中から指定してください: {text!r}")
        if field.type == "number":
            try:
                float(text)
            except ValueError:
                raise _invalid_param(f"「{field.label}」には数値を入力してください: {text!r}") from None
        params[field.name] = text
    return params


def _check_fields(fields: Sequence[ParamField]) -> None:
    seen: set[str] = set()
    for field in fields:
        if not _PARAM_NAME_RE.fullmatch(field.name):
            raise ValueError(f"入力項目の name には英数字と _ のみ使えます: {field.name!r}")
        if field.name in seen:
            raise ValueError(f"入力項目の name が重複しています: {field.name!r}")
        seen.add(field.name)
        if field.type == "select":
            if not field.choices:
                raise ValueError(f"select の入力項目には choices が必要です: {field.name!r}")
            if field.default and field.default not in field.choices:
                raise ValueError(f"default が choices に含まれていません: {field.name!r}")


def load_customization(spec: str | None = None) -> WebCustomization:
    """``module:ClassName`` 形式の指定からカスタマイズを読み込む。

    ``spec`` を省略した場合は環境変数 ``CHILMAI_WEB_CUSTOMIZATION`` を使い、
    どちらも未指定の場合は標準の ``WebCustomization`` を返す。
    指定に誤りがある場合は、標準の挙動に切り替えずに例外を送出する。
    """
    if spec is None:
        spec = os.getenv(ENV_VAR, "")
    spec = spec.strip()
    if not spec:
        return WebCustomization()

    module_name, sep, class_name = spec.partition(":")
    if not sep or not module_name or not class_name:
        raise ValueError(f"{ENV_VAR} は module:ClassName の形式で指定してください: {spec!r}")
    module = importlib.import_module(module_name)
    cls = getattr(module, class_name, None)
    if not (isinstance(cls, type) and issubclass(cls, WebCustomization)):
        raise TypeError(f"{spec!r} は WebCustomization のサブクラスではありません。")
    customization = cls()
    _check_fields(customization.fields)
    return customization
