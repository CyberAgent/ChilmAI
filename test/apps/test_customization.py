from __future__ import annotations

import base64
import io
import json
import re
from typing import ClassVar

import openpyxl
import pandas as pd
import pytest

fastapi = pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

import apps.api.main as api_main  # noqa: E402
from apps.api.customization import (  # noqa: E402
    ParamField,
    WebCustomization,
    load_customization,
    parse_params,
)
from chilmai.generic.config import DEFAULT_CONFIG  # noqa: E402
from chilmai.generic.error_codes import ChilmError, ErrorCode  # noqa: E402
from chilmai.generic.parser import InputParser  # noqa: E402
from chilmai.generic.service import MatchingService  # noqa: E402
from examples.sample_web_customization import SampleWebCustomization  # noqa: E402

SAMPLE_SPEC = "examples.sample_web_customization:SampleWebCustomization"


def _children_csv() -> bytes:
    return (
        "child_id,household_id,age,score_1,pref_1,pref_2,enrolled_daycare_id,sibling_pattern\n"
        "1,10,1,100,100,101,,\n"
        "2,11,1,80,100,101,,\n"
        "3,12,1,60,101,100,,\n"
    ).encode("utf-8")


def _daycares_csv() -> bytes:
    return (
        "daycare_id,daycare_name,capacity_age0,capacity_age1,capacity_age2,capacity_age3,capacity_age4,capacity_age5\n"
        "100,A,0,1,0,0,0,0\n"
        "101,B,0,1,0,0,0,0\n"
    ).encode("utf-8")


def _children_xlsx_with_title_row() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["令和9年度 申込者一覧"])
    for line in _children_csv().decode("utf-8").splitlines():
        ws.append(line.split(","))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload_files(children: tuple[str, bytes, str] | None = None) -> dict:
    return {
        "children_file": children or ("children.csv", io.BytesIO(_children_csv()), "text/csv"),
        "daycares_file": ("daycares.csv", io.BytesIO(_daycares_csv()), "text/csv"),
    }


@pytest.fixture
def sample_customization(monkeypatch):
    customization = load_customization(SAMPLE_SPEC)
    monkeypatch.setattr(api_main, "customization", customization)
    return customization


# --- parse_params ---------------------------------------------------------

_FIELDS = (
    ParamField(name="year", label="年度", type="number", required=True),
    ParamField(name="month", label="入所月", type="select", choices=("4", "5"), default="4"),
    ParamField(name="note", label="備考"),
    ParamField(name="flag", label="フラグ", type="checkbox", default="true"),
)


def test_parse_params_fills_defaults_and_normalizes_values():
    params = parse_params(_FIELDS, {"year": " 2027 "})
    assert params == {"year": "2027", "month": "4", "note": "", "flag": "true"}


@pytest.mark.parametrize(
    ("raw", "message"),
    [
        ({}, "「年度」を入力してください"),
        ({"year": ""}, "「年度」を入力してください"),
        ({"year": "令和9"}, "「年度」には半角の数値を入力してください"),
        ({"year": "nan"}, "「年度」には半角の数値を入力してください"),
        ({"year": "inf"}, "「年度」には半角の数値を入力してください"),
        ({"year": "２０２７"}, "「年度」には半角の数値を入力してください"),
        ({"year": "2027", "month": "6"}, "「入所月」には選択肢の中から指定してください"),
        ({"year": "2027", "unknown": "x"}, "定義されていない入力項目"),
    ],
)
def test_parse_params_rejects_invalid_values(raw, message):
    with pytest.raises(ChilmError, match=message) as exc_info:
        parse_params(_FIELDS, raw)
    assert exc_info.value.code == ErrorCode.INVALID_INPUT_PARAMETER


@pytest.mark.parametrize("value", ["2027", "-1", "0.5", ".5", "+3"])
def test_parse_params_accepts_ascii_numbers(value):
    assert parse_params(_FIELDS, {"year": value})["year"] == value


def test_parse_params_keeps_explicit_empty_value_instead_of_default():
    # 任意の select で「—」を選んだ場合は、既定値で上書きしない。
    assert parse_params(_FIELDS, {"year": "2027", "month": ""})["month"] == ""


@pytest.mark.parametrize(("value", "expected"), [("true", "true"), ("on", "true"), ("false", "false")])
def test_parse_params_checkbox_values(value, expected):
    assert parse_params(_FIELDS, {"year": "2027", "flag": value})["flag"] == expected


# --- load_customization ---------------------------------------------------


def test_load_customization_defaults_to_base_class(monkeypatch):
    monkeypatch.delenv("CHILMAI_WEB_CUSTOMIZATION", raising=False)
    assert type(load_customization()) is WebCustomization


def test_load_customization_reads_env_var(monkeypatch):
    monkeypatch.setenv("CHILMAI_WEB_CUSTOMIZATION", SAMPLE_SPEC)
    assert isinstance(load_customization(), SampleWebCustomization)


@pytest.mark.parametrize(
    ("spec", "error"),
    [
        ("examples.sample_web_customization", ValueError),
        ("examples.sample_web_customization:Missing", TypeError),
        ("examples.sample_web_customization:TitleRowParser", TypeError),
        ("no_such_module_for_chilmai:Foo", ModuleNotFoundError),
    ],
)
def test_load_customization_rejects_invalid_spec(spec, error):
    with pytest.raises(error):
        load_customization(spec)


class _DuplicateFields(WebCustomization):
    fields: ClassVar = (ParamField(name="a", label="A"), ParamField(name="a", label="A2"))


class _BadName(WebCustomization):
    fields: ClassVar = (ParamField(name="年度", label="年度"),)


class _SelectWithoutChoices(WebCustomization):
    fields: ClassVar = (ParamField(name="a", label="A", type="select"),)


class _UnknownType(WebCustomization):
    fields: ClassVar = (ParamField(name="a", label="A", type="date"),)  # type: ignore[arg-type]


class _CheckboxDefaultNotLowercase(WebCustomization):
    fields: ClassVar = (ParamField(name="a", label="A", type="checkbox", default="True"),)


class _NumberDefaultNotNumeric(WebCustomization):
    fields: ClassVar = (ParamField(name="a", label="A", type="number", default="abc"),)


@pytest.mark.parametrize(
    "cls_name",
    [
        "_DuplicateFields",
        "_BadName",
        "_SelectWithoutChoices",
        "_UnknownType",
        "_CheckboxDefaultNotLowercase",
        "_NumberDefaultNotNumeric",
    ],
)
def test_load_customization_rejects_invalid_fields(cls_name):
    with pytest.raises(ValueError):
        load_customization(f"{__name__}:{cls_name}")


# --- InputParser の読み込み設定 ------------------------------------------


class _HeaderRow2Parser(InputParser):
    SHEET_NAME = "申込"
    HEADER_ROW = 2


def _xlsx_with_sheets() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "表紙"
    wb.active.append(["これは表紙です"])
    ws = wb.create_sheet("申込")
    ws.append(["タイトル"])
    ws.append(["child_id", "age"])
    ws.append(["1", "2"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parser_subclass_reads_configured_sheet_and_header_row():
    parser = _HeaderRow2Parser()
    raw = parser.read_raw(file_bytes=_xlsx_with_sheets(), file_format="xlsx")
    assert list(raw.columns) == ["child_id", "age"]
    assert raw.to_dict(orient="records") == [{"child_id": "1", "age": "2"}]
    assert parser.read_columns(io.BytesIO(_xlsx_with_sheets()), "xlsx") == ["child_id", "age"]


def test_parser_header_row_applies_to_csv():
    csv_bytes = "タイトル行\nchild_id,age\n1,2\n".encode("utf-8")
    assert _HeaderRow2Parser.read_columns(io.BytesIO(csv_bytes), "csv") == ["child_id", "age"]


def test_matching_service_uses_given_parser_with_file_kind():
    kinds: list[str | None] = []

    class RecordingParser(InputParser):
        @classmethod
        def read_table(cls, source, file_format, *, kind=None, nrows=None) -> pd.DataFrame:
            kinds.append(kind)
            return super().read_table(source, file_format, kind=kind, nrows=nrows)

    service = MatchingService(parser=RecordingParser())
    service.match(
        children_file_bytes=_children_csv(),
        children_file_format="csv",
        daycares_file_bytes=_daycares_csv(),
        daycares_file_format="csv",
        mapping=DEFAULT_CONFIG,
        solver_config={"max_time_seconds": 3},
    )
    # read_raw（結果ファイル作成用の申込者データの再読込）も kind 付きで呼ばれる。
    assert kinds == ["children", "daycares", "children"]


# --- リファレンスアプリ ---------------------------------------------------


def test_top_page_without_customization_has_no_param_inputs():
    resp = TestClient(api_main.app).get("/")
    assert resp.status_code == 200
    assert 'name="param__' not in resp.text


def test_top_page_renders_param_inputs(sample_customization):
    resp = TestClient(api_main.app).get("/")
    assert resp.status_code == 200
    assert 'name="param__fiscal_year"' in resp.text
    assert 'name="param__month"' in resp.text
    assert re.search(r'name="param__add_remarks" value="true"\s+checked', resp.text)


def test_htmx_match_uses_customized_excel_and_filename(sample_customization):
    resp = TestClient(api_main.app).post(
        "/htmx/match",
        data={"param__fiscal_year": "2027", "param__month": "5", "param__add_remarks": "true"},
        files=_upload_files(),
    )
    assert resp.status_code == 200
    assert resp.headers["X-Result-Mode"] == "match"
    assert 'data-download-filename="選考結果_2027年度5月.xlsx"' in resp.text

    b64 = re.search(r'id="matching-result-excel"[^>]*>([^<]+)<', resp.text).group(1).strip()
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    assert wb.sheetnames == ["result", "施設別", "meta"]
    assert wb["施設別"]["A1"].value == "2027年度 5月入所 選考結果（施設別）"
    result_header = [cell.value for cell in wb["result"][1]]
    assert result_header[-1] == "備考"
    meta_rows = {row[0].value: row[1].value for row in wb["meta"].iter_rows(min_row=2)}
    assert meta_rows["param.fiscal_year"] == "2027"
    assert meta_rows["param.add_remarks"] == "true"


def test_htmx_match_unchecked_checkbox_disables_preprocessor(sample_customization):
    resp = TestClient(api_main.app).post(
        "/htmx/match",
        data={"param__fiscal_year": "2027", "param__month": "4"},
        files=_upload_files(),
    )
    b64 = re.search(r'id="matching-result-excel"[^>]*>([^<]+)<', resp.text).group(1).strip()
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    assert "備考" not in [cell.value for cell in wb["result"][1]]


def test_htmx_validate_reports_invalid_param(sample_customization):
    resp = TestClient(api_main.app).post(
        "/htmx/validate",
        data={"param__fiscal_year": "2027", "param__month": "13"},
        files=_upload_files(),
    )
    assert resp.headers["X-Result-Mode"] == "validate-error"
    assert "E108" in resp.text


def test_htmx_validate_reads_children_excel_with_customized_parser(sample_customization):
    xlsx = _children_xlsx_with_title_row()
    resp = TestClient(api_main.app).post(
        "/htmx/validate",
        data={"param__fiscal_year": "2027", "param__month": "4"},
        files=_upload_files(
            ("children.xlsx", io.BytesIO(xlsx), "application/vnd.ms-excel"),
        ),
    )
    assert resp.headers["X-Validation-Status"] == "valid"


def test_analyze_columns_uses_customized_parser(sample_customization):
    xlsx = _children_xlsx_with_title_row()
    resp = TestClient(api_main.app).post(
        "/htmx/analyze-columns",
        files={"children_file": ("children.xlsx", io.BytesIO(xlsx), "application/vnd.ms-excel")},
    )
    assert resp.status_code == 200
    assert "children.xlsx" in resp.text
    assert "解析に失敗" not in resp.text


def test_api_match_accepts_params_json(sample_customization):
    resp = TestClient(api_main.app).post(
        "/match",
        data={"params": json.dumps({"fiscal_year": "2027", "month": "4"})},
        files=_upload_files(),
    )
    assert resp.status_code == 200
    assert resp.json()["meta"]["algorithm"] == "cp_use_transfer"


@pytest.mark.parametrize("params", ["not json", "[1, 2]", json.dumps({"fiscal_year": "2027", "x": "1"})])
def test_api_validate_rejects_invalid_params(sample_customization, params):
    resp = TestClient(api_main.app).post("/validate", data={"params": params}, files=_upload_files())
    assert resp.status_code == 422
    error = resp.json()["detail"]["errors"][0]
    assert error["code"] == ErrorCode.INVALID_INPUT_PARAMETER
    assert error["type"] == "config"


def test_api_validate_without_customization_ignores_missing_params():
    resp = TestClient(api_main.app).post("/validate", files=_upload_files())
    assert resp.status_code == 200
    assert resp.json()["is_valid"] is True


def test_customized_workbook_formulas_are_neutralized(monkeypatch):
    class FormulaCustomization(WebCustomization):
        def build_result_workbook(self, result, meta, params):
            wb = openpyxl.Workbook()
            wb.active.append(["=1+1"])
            wb.create_sheet("extra").append(['=HYPERLINK("http://evil.example.com")'])
            return wb

    monkeypatch.setattr(api_main, "customization", FormulaCustomization())
    b64 = api_main._build_matching_excel({"output_columns": [], "output_rows": []})
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f"


class _BrokenCustomization(WebCustomization):
    def __init__(self, broken: str) -> None:
        self.broken = broken

    def create_preprocessor(self, params):
        if self.broken == "create_preprocessor":
            raise RuntimeError("boom")
        return None

    def build_result_workbook(self, result, meta, params):
        if self.broken == "build_result_workbook":
            raise KeyError("missing_key")
        return super().build_result_workbook(result, meta, params)


@pytest.mark.parametrize("broken", ["create_preprocessor", "build_result_workbook"])
def test_htmx_match_reports_customization_error(monkeypatch, broken):
    monkeypatch.setattr(api_main, "customization", _BrokenCustomization(broken))
    resp = TestClient(api_main.app).post("/htmx/match", files=_upload_files())
    assert resp.status_code == 200
    assert resp.headers["X-Result-Mode"] == "validate-error"
    assert "E109" in resp.text
    assert broken in resp.text


def test_api_validate_reports_customization_error(monkeypatch):
    monkeypatch.setattr(api_main, "customization", _BrokenCustomization("create_preprocessor"))
    resp = TestClient(api_main.app).post("/validate", files=_upload_files())
    assert resp.status_code == 422
    error = resp.json()["detail"]["errors"][0]
    assert error["code"] == ErrorCode.CUSTOMIZATION_FAILED
    assert error["type"] == "config"
