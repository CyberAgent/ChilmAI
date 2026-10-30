"""リファレンスアプリ（Web UI・サーバ API）のカスタマイズのサンプル。

``WebCustomization`` を継承し、次の 4 か所を差し替える例を示す。

- 追加の入力項目: 年度・入所月・結果の備考列の有無を画面に追加する。
- 前処理フック: 入力項目の値に応じて、結果ファイルに備考列を追加する。
- 入力ファイルの読み込み: 申込者データの Excel だけ、1 行目の帳票タイトルを
  読み飛ばし、2 行目を列名として読む。
- 結果 Excel: 標準の ``result``・``meta`` シートに、見出し付きの施設別シートを追加する。

使い方（リポジトリのルートで実行する）::

    CHILMAI_WEB_CUSTOMIZATION=examples.sample_web_customization:SampleWebCustomization \\
        uv run uvicorn apps.api.main:app --host 127.0.0.1 --port 8501
"""

from __future__ import annotations

from collections.abc import Mapping
from typing import IO, Any

import openpyxl
import pandas as pd
from openpyxl.styles import Font

from apps.api.customization import ParamField, WebCustomization, build_default_workbook
from chilmai.generic.parser import FileFormat, FileKind, InputParser
from chilmai.generic.preprocessor import BasePreprocessor


class TitleRowParser(InputParser):
    """申込者データの Excel だけ、2 行目を列名として読む。

    対象のシートと列名の行がすべてのファイルで同じ場合は、``read_table`` を
    オーバーライドせずに、クラス属性を上書きするだけでよい::

        class MyParser(InputParser):
            SHEET_NAME = "申込一覧"
            HEADER_ROW = 2
    """

    @classmethod
    def read_table(
        cls,
        source: IO[bytes],
        file_format: FileFormat,
        *,
        kind: FileKind | None = None,
        nrows: int | None = None,
    ) -> pd.DataFrame:
        if kind == "children" and file_format == "excel":
            return pd.read_excel(source, sheet_name=0, dtype=str, header=1, nrows=nrows, engine="calamine")
        return super().read_table(source, file_format, kind=kind, nrows=nrows)


class RemarksPreprocessor(BasePreprocessor):
    """結果ファイルの末尾に、内定の有無を表す備考列を追加する。"""

    def __init__(self, label: str) -> None:
        self.label = label

    def transform_output(self, df: pd.DataFrame, result: dict[str, Any]) -> pd.DataFrame:
        out = df.copy()
        # 結果列の列名は設定で変えられるため、位置で参照する。
        # transform_output の時点では、結果の保育所ID・保育所名の列がこの順で末尾にある。
        result_id_col, _result_name_col = out.columns[-2:]
        out["備考"] = out[result_id_col].map(lambda v: f"{self.label} 内定" if v else "保留")
        return out


class SampleWebCustomization(WebCustomization):
    """年度・入所月を入力し、施設別シート付きの結果 Excel を作るカスタマイズ。"""

    fields = (
        ParamField(
            name="fiscal_year",
            label="年度",
            type="number",
            default="2027",
            required=True,
            help_text="結果 Excel の見出しとファイル名に使います。",
        ),
        ParamField(
            name="month",
            label="入所月",
            type="select",
            choices=tuple(str(m) for m in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)),
            default="4",
            required=True,
        ),
        ParamField(
            name="add_remarks",
            label="結果ファイルに備考列を追加する",
            type="checkbox",
            default="true",
        ),
    )

    def create_preprocessor(self, params: Mapping[str, str]) -> BasePreprocessor | None:
        if params["add_remarks"] != "true":
            return None
        return RemarksPreprocessor(label=f"{params['month']}月入所")

    def create_parser(self) -> InputParser:
        return TitleRowParser()

    def build_result_workbook(
        self,
        result: dict[str, Any],
        meta: Mapping[str, str],
        params: Mapping[str, str],
    ) -> openpyxl.Workbook:
        wb = build_default_workbook(result, meta, params)

        ws = wb.create_sheet(title="施設別", index=1)
        ws.append([f"{params['fiscal_year']}年度 {params['month']}月入所 選考結果（施設別）"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(["保育所ID", "保育所名", "内定数"])
        for cell in ws[3]:
            cell.font = Font(bold=True)

        counts: dict[str, int] = {}
        for daycare_id in result.get("matching_result_dict", {}).values():
            if daycare_id is not None:
                counts[daycare_id] = counts.get(daycare_id, 0) + 1
        names: dict[str, str] = result.get("daycare_name_dict", {})
        for daycare_id in sorted(counts):
            ws.append([daycare_id, names.get(daycare_id, ""), counts[daycare_id]])
        return wb

    def result_filename(self, params: Mapping[str, str]) -> str:
        return f"選考結果_{params['fiscal_year']}年度{params['month']}月.xlsx"
